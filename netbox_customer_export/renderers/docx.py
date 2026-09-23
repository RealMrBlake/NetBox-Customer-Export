import zipfile
from io import BytesIO

from . import MissingDependency

DOTX_TYPE = b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
DOCX_TYPE = b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"


def _svg_to_png(svg):
    try:
        import cairosvg
    except (ImportError, OSError):
        return None
    try:
        return cairosvg.svg2png(bytestring=svg.encode("utf-8"), scale=2)
    except Exception:  # noqa: BLE001 – dann greift das Pillow-PNG
        return None


def _load_template(path):
    """Öffnet eine .docx/.dotx-Vorlage und leert deren Textkörper (Formatvorlagen, Kopf-/Fußzeile bleiben)."""
    from docx import Document as Docx

    buf = BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(DOTX_TYPE, DOCX_TYPE)  # .dotx -> .docx, sonst lehnt python-docx ab
            dst.writestr(item, data)
    buf.seek(0)
    d = Docx(buf)
    body = d.element.body
    for child in list(body):
        if not child.tag.endswith("}sectPr"):
            body.remove(child)
    return d


def render_docx(doc):
    try:
        from docx import Document as Docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as e:
        raise MissingDependency(f"Word-Export benötigt python-docx (pip install python-docx cairosvg): {e}")

    b = doc.branding
    primary = RGBColor.from_string(b["primary_color"])
    template = b.get("docx_template")

    if template:
        try:
            d = _load_template(template)
        except (OSError, KeyError, zipfile.BadZipFile) as e:
            raise MissingDependency(f"Word-Vorlage {template!r} konnte nicht geladen werden: {e}")
    else:
        # Ohne Vorlage: Standarddokument mit Farben/Schrift aus dem Branding
        d = Docx()
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)  # A4 statt US-Letter aus der python-docx-Vorlage
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin = sec.bottom_margin = Cm(2)
        d.styles["Normal"].font.name = b["font_family"]
        d.styles["Normal"].font.size = Pt(b["font_size"])
        for level, color in ((1, b["primary_color"]), (2, b["primary_color"]), (3, b["accent_color"]), (4, b["accent_color"])):
            style = d.styles[f"Heading {level}"]
            style.font.name = b["font_family"]
            style.font.color.rgb = RGBColor.from_string(color)
    sec = d.sections[0]

    style_names = {s.name for s in d.styles}
    table_style = b.get("docx_table_style") if b.get("docx_table_style") in style_names else None

    def shade(cell, color):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), color)
        tc_pr.append(shd)

    def set_style(table):
        if table_style:
            table.style = table_style
        elif "Table Grid" in style_names:
            table.style = "Table Grid"

    # Deckblatt
    meta = doc.meta
    if b.get("logo_path"):
        d.add_picture(b["logo_path"], width=Cm(5))
    for _ in range(6):
        d.add_paragraph()
    title = d.add_paragraph()
    run = title.add_run(doc.title)
    run.bold, run.font.size = True, Pt(28)
    run.font.color.rgb = primary
    if doc.subtitle:
        d.add_paragraph().add_run(doc.subtitle).font.size = Pt(16)
    d.add_paragraph()
    info = d.add_paragraph()
    for label, value in (("Stand", meta.get("date")), ("Erstellt von", meta.get("author")), ("", b.get("company_name"))):
        if value:
            info.add_run(f"{label + ': ' if label else ''}{value}\n")
    d.add_page_break()

    # Inhaltsverzeichnis (wird beim Öffnen in Word aktualisiert)
    d.add_heading("Inhaltsverzeichnis", level=1)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
    d.add_paragraph()._p.append(fld)
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    d.settings.element.append(upd)
    d.add_page_break()

    from docx.enum.section import WD_ORIENT, WD_SECTION

    portrait = d.sections[-1]
    portrait_size = (portrait.page_width, portrait.page_height)
    landscape_open = False

    def new_section(landscape):
        s = d.add_section(WD_SECTION.NEW_PAGE)
        w, h = portrait_size
        s.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = (h, w) if landscape else (w, h)
        return s

    for blk in doc.blocks:
        if landscape_open and blk.kind != "landscape":
            # erst beim nächsten Inhalt zurück ins Hochformat – vermeidet leere Seiten
            new_section(landscape=False)
            landscape_open = False
        if blk.kind == "landscape":
            s = new_section(landscape=True)
            landscape_open = True
            for h in blk.rows:
                d.add_heading(h.text, level=h.level)
            png = blk.images[0]["png"]
            from PIL import Image

            px_w, px_h = Image.open(BytesIO(png)).size
            avail_w = s.page_width - s.left_margin - s.right_margin
            avail_h = s.page_height - s.top_margin - s.bottom_margin - Cm(1.4 * len(blk.rows) + 0.8)
            para = d.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if px_w / px_h >= avail_w / avail_h:
                para.add_run().add_picture(BytesIO(png), width=avail_w)
            else:
                para.add_run().add_picture(BytesIO(png), height=avail_h)
        elif blk.kind == "heading":
            d.add_heading(blk.text, level=min(blk.level, 4))
        elif blk.kind == "paragraph":
            d.add_paragraph(blk.text)
        elif blk.kind == "pagebreak":
            d.add_page_break()
        elif blk.kind in ("table", "kv"):
            cols = len(blk.headers) if blk.kind == "table" else 2
            t = d.add_table(rows=0, cols=cols)
            set_style(t)
            if blk.kind == "table":
                cells = t.add_row().cells
                for c, h in zip(cells, blk.headers):
                    c.text = h
                    r = c.paragraphs[0].runs[0]
                    r.bold = True
                    if not table_style:  # eine Tabellenformatvorlage bringt eigene Kopfzeilenfarben mit
                        r.font.color.rgb = RGBColor.from_string(b["header_text_color"])
                        shade(c, b["primary_color"])
                # Kopfzeile auf jeder Seite wiederholen
                tr_pr = t.rows[0]._tr.get_or_add_trPr()
                hdr = OxmlElement("w:tblHeader")
                hdr.set(qn("w:val"), "true")
                tr_pr.append(hdr)
            for row in blk.rows:
                cells = t.add_row().cells
                for i, (c, v) in enumerate(zip(cells, row)):
                    c.text = v
                    for p in c.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(max(b["font_size"] - 1, 7))
                    if blk.kind == "kv" and i == 0:
                        c.paragraphs[0].runs[0].bold = True
                        if not table_style:
                            shade(c, b["label_bg_color"])
            d.add_paragraph()
        elif blk.kind == "images":
            t = d.add_table(rows=2, cols=len(blk.images))
            for i, img in enumerate(blk.images):
                png = (_svg_to_png(img["svg"]) if img["svg"] else None) or img["png"]
                cell = t.cell(0, i).paragraphs[0]
                cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if png:
                    # Topologie: volle Breite; Racks: 7 cm (Höhe skaliert mit der Rackgröße)
                    cell.add_run().add_picture(BytesIO(png), width=Cm(16.5 if blk.level == 1 else 7))
                else:
                    cell.add_run("(Rack-Grafik nicht verfügbar)")
                cap = t.cell(1, i).paragraphs[0]
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if img["caption"]:
                    cap.add_run(img["caption"]).italic = True
            d.add_paragraph()

    # Fußzeile nur setzen, wenn die Vorlage keine eigene mitbringt
    footer = sec.footer
    if not any(p.text.strip() for p in footer.paragraphs):
        text = " · ".join(x for x in (b.get("company_name"), b.get("footer_text"), doc.title) if x)
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.add_run(text).font.size = Pt(8)

    out = BytesIO()
    d.save(out)
    return out.getvalue()
